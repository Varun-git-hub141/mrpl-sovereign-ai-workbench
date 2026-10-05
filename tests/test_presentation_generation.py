import os
import pytest
from pathlib import Path
from pptx import Presentation
from fastapi.testclient import TestClient

from backend.app import app
from backend.services.deliverable_generator import DeliverableGenerator
from backend.services.intent_router import IntentRouter, UserIntent, OutputFormat


@pytest.fixture
def client():
    return TestClient(app)


def test_extract_topic_title_various_queries():
    # Test 1: Explicit entity tag
    assert DeliverableGenerator.extract_topic_title("some query", "P-201") == "P-201"
    assert DeliverableGenerator.extract_topic_title("some query", "V-201") == "V-201"
    assert DeliverableGenerator.extract_topic_title("some query", "CDU-P-101") == "CDU-P-101"

    # Test 2: Natural language queries without entity tag
    t1 = DeliverableGenerator.extract_topic_title("Create a presentation on boiler feed water treatment")
    assert t1 == "Boiler Feed Water Treatment"

    t2 = DeliverableGenerator.extract_topic_title("Prepare a slide deck about flare gas recovery systems")
    assert t2 == "Flare Gas Recovery Systems"

    t3 = DeliverableGenerator.extract_topic_title("Make a presentation on refinery safety protocols in 5 slides")
    assert t3 == "Refinery Safety Protocols"

    t4 = DeliverableGenerator.extract_topic_title("Presentation regarding crude distillation unit operations")
    assert t4 == "Crude Distillation Unit Operations"

    t5 = DeliverableGenerator.extract_topic_title("Give me slides on heat exchanger maintenance")
    assert t5 == "Heat Exchanger Maintenance"


def test_generate_pptx_p201():
    result = DeliverableGenerator.generate_pptx(
        entity_tag="P-201",
        query="Create a presentation on P-201 asset integrity",
        retrieved_chunks=[],
        calc_result={"status": "SUCCESS", "formula_id": "CR-001", "result": 0.30, "result_unit": "mm/year"},
    )
    assert result["deliverable_type"] == "PPTX"
    assert "P-201" in result["title"] or "P_201" in result["filename"]
    assert os.path.exists(result["filepath"])

    prs = Presentation(result["filepath"])
    assert len(prs.slides) >= 4

    # Slide 1 Check
    s1_text = "".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "MANGALORE REFINERY" in s1_text
    assert "P-201" in s1_text

    # Slide 2 Check
    s2_text = "".join(shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "Crude Feed Booster Pump" in s2_text


def test_generate_pptx_other_topic_boiler_water():
    topic = "Boiler Feed Water Treatment"
    query = "Create a presentation on boiler feed water treatment"
    summary = """
    Boiler feed water treatment is critical to prevent scale formation and corrosion in refinery steam boilers.
    • Demineralization and reverse osmosis ensure high water purity.
    • Oxygen scavenging using hydrazine or DEHA prevents pitting corrosion.
    • Phosphate treatment maintains alkaline pH and prevents scale.
    • Continuous blowdown control manages dissolved solids (TDS).
    • ASME guidelines dictate strict silica and conductivity limits.
    """

    result = DeliverableGenerator.generate_pptx(
        entity_tag=None,
        query=query,
        retrieved_chunks=[],
        summary_text=summary,
    )

    assert result["deliverable_type"] == "PPTX"
    assert "Boiler_Feed_Water_Treatment" in result["filename"]
    assert "Boiler Feed Water Treatment" in result["title"]
    assert os.path.exists(result["filepath"])

    prs = Presentation(result["filepath"])
    assert len(prs.slides) >= 4

    # Slide 1 Check
    s1_text = "".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "MANGALORE REFINERY" in s1_text
    assert "Boiler Feed Water Treatment" in s1_text
    assert "P-201" not in s1_text  # MUST NOT contain P-201

    # Slide 2 Check
    s2_text = "".join(shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "Boiler Feed Water Treatment" in s2_text
    assert "Demineralization" in s2_text or "Oxygen scavenging" in s2_text or "Boiler Feed Water" in s2_text
    assert "Crude Feed Booster Pump" not in s2_text  # MUST NOT contain P-201 pump


def test_generate_pptx_other_topic_flare_system():
    query = "Make a presentation about flare gas recovery systems"
    result = DeliverableGenerator.generate_pptx(
        entity_tag=None,
        query=query,
        retrieved_chunks=[],
        summary_text="Flare gas recovery minimizes refinery hydrocarbon emissions and recovers fuel gas.",
    )

    assert result["deliverable_type"] == "PPTX"
    assert "Flare Gas Recovery" in result["title"]
    assert "P-201" not in result["title"]

    prs = Presentation(result["filepath"])
    s1_text = "".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Flare Gas Recovery" in s1_text
    assert "P-201" not in s1_text


def test_api_e2e_presentation_generation_other_topic(client):
    query = "Create a presentation on refinery safety protocols"

    # 1. Test Intent Detection
    intent = IntentRouter.detect_intent(query=query)
    assert intent.intent == UserIntent.PRESENTATION_GENERATION
    assert intent.output_format == OutputFormat.PPTX

    # 2. Test Route
    route_res = client.post("/api/route", json={"query": query})
    assert route_res.status_code == 200
    route_data = route_res.json()
    assert route_data["detected_intent"] == "PRESENTATION_GENERATION"
    assert route_data["inferred_output_format"] == "PPTX"

    # 3. Test Execute
    exec_res = client.post("/api/execute", json={
        "query": query,
        "execution_plan": route_data["execution_plan"],
        "allow_unverified_fallback": True,
    })
    assert exec_res.status_code == 200
    exec_data = exec_res.json()

    deliv = exec_data.get("generated_deliverable")
    assert deliv is not None
    assert deliv["deliverable_type"] == "PPTX"
    assert "Refinery Safety Protocols" in deliv["title"]
    assert "P-201" not in deliv["title"]
    assert os.path.exists(deliv["filepath"])


def test_generate_pptx_sarojini_naidu_biographical_paragraphs():
    query = "Create a presentation on Sarojini Naidu"
    # Text with conversational disclaimer that must be stripped out and parsed into distinct biographical slides
    raw_text = """I'm sorry, but I can't create a presentation on Sarojini Naidu as I don't have access to her personal information or experiences. However, I can provide you with some general information about Sarojini Naidu and her contributions to Indian history and culture.

Sarojini Naidu was a prominent Indian poet, freedom fighter, and political leader. She was born on February 13, 1879, in Hyderabad, India. She was famously known as the Nightingale of India (Bharat Kokila).

Naidu was a passionate advocate for women's rights and education. She founded the Women's Indian Association in 1917, which aimed to promote women's education and empowerment. She also advocated for women's suffrage and equal political rights across India.

Naidu was also a prominent figure in the Indian independence movement. She was arrested multiple times for her involvement in protests and the Salt March with Mahatma Gandhi. She became the first Indian woman to serve as President of the Indian National Congress in 1925.

Naidu was known for her lyrical poetry, which often focused on themes of love, patriotism, freedom, and nature. Her famous collections include The Golden Threshold and The Bird of Time.

Naidu passed away on March 2, 1949, while serving as the Governor of the United Provinces, becoming the first woman governor in independent India."""

    result = DeliverableGenerator.generate_pptx(
        entity_tag=None,
        query=query,
        retrieved_chunks=[],
        summary_text=raw_text,
    )

    assert result["deliverable_type"] == "PPTX"
    assert "Sarojini Naidu" in result["title"]
    assert "P-201" not in result["title"]
    assert os.path.exists(result["filepath"])

    prs = Presentation(result["filepath"])
    assert len(prs.slides) >= 5

    # Check Slide 1 (Title)
    s1_text = "".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Sarojini Naidu" in s1_text
    assert "P-201" not in s1_text

    # Check that apology/refusal is NOT anywhere in the presentation
    all_text = " ".join("".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in prs.slides)
    assert "I'm sorry" not in all_text
    assert "personal information or experiences" not in all_text

    # Check that real biographical content is present in the slides
    assert "Nightingale of India" in all_text or "freedom fighter" in all_text
    assert "Women's Indian Association" in all_text or "women's rights" in all_text
    assert "Indian National Congress" in all_text or "independence movement" in all_text
    assert "Core Subject Scope" not in all_text  # Fallback template avoided
